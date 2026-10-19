import asyncio
import csv
import html
import io
import logging
import os
import re
import time
from datetime import datetime, timedelta, timezone
from typing import Optional

from aiogram import Bot, Dispatcher, F, Router, types
from aiogram.client.default import DefaultBotProperties
from aiogram.enums import ParseMode
from aiogram.exceptions import TelegramBadRequest, TelegramForbiddenError, TelegramRetryAfter
from aiogram.filters import Command, CommandStart, StateFilter
from aiogram.fsm.context import FSMContext
from aiogram.fsm.state import State, StatesGroup
from aiogram.fsm.storage.memory import MemoryStorage
from aiogram.types import (
    BotCommand, BotCommandScopeChat, BufferedInputFile, InlineKeyboardButton,
    InlineKeyboardMarkup, KeyboardButton, ReplyKeyboardMarkup, ReplyKeyboardRemove,
)
from aiogram.utils.keyboard import InlineKeyboardBuilder
from dotenv import load_dotenv
from sqlalchemy import (
    BigInteger, Boolean, DateTime, Integer, String, Text, UniqueConstraint,
    desc, func, inspect, select,
)
from sqlalchemy import text as sa_text
from sqlalchemy.exc import IntegrityError
from sqlalchemy.ext.asyncio import async_sessionmaker, create_async_engine
from sqlalchemy.orm import DeclarativeBase, Mapped, mapped_column

load_dotenv(override=True)
TOKEN = os.getenv("BOT_TOKEN", "").strip()
ADMIN_IDS = {int(x) for x in re.split(r"[,\s]+", os.getenv("ADMIN_IDS", "").strip()) if x.isdigit()}
DATABASE_URL = os.getenv("DATABASE_URL", "sqlite+aiosqlite:///feedback.db")
TZ = timezone(timedelta(hours=5))
MAX_LEN, MIN_LEN, PAGE_SIZE, COOLDOWN_SEC = 3000, 5, 8, 15
_last_sent: dict[int, float] = {}

LANGS = {"uz": "🇺🇿 O'zbekcha", "ru": "🇷🇺 Русский", "en": "🇬🇧 English"}
DEFAULT_LANG = "uz"

def _t(uz, ru, en): return {"uz": uz, "ru": ru, "en": en}
T = {
"choose_lang": _t("🌐 Tilni tanlang", "🌐 Выберите язык", "🌐 Choose language"),
"lang_set": _t("✅ Til o'zgartirildi", "✅ Язык изменён", "✅ Language changed"),
"cancelled": _t("❌ Bekor qilindi", "❌ Отменено", "❌ Cancelled"),
"cancel": _t("❌ Bekor qilish", "❌ Отмена", "❌ Cancel"),
"back": _t("🔙 Orqaga", "🔙 Назад", "🔙 Back"),
"only_text": _t("Iltimos, matn yuboring.", "Пожалуйста, отправьте текст.", "Please send text."),
"welcome": _t("Assalomu alaykum! Ism-familiyangiz va guruhingizni yozing.\nMasalan: Aliyev Vali, 12-guruh", "Здравствуйте! Напишите имя, фамилию и группу.\nНапример: Алиев Вали, группа 12", "Hello! Enter your full name and group.\nExample: Ali Valiyev, group 12"),
"hello": _t("Assalomu alaykum, <b>{name}</b>!", "Здравствуйте, <b>{name}</b>!", "Hello, <b>{name}</b>!"),
"need_name": _t("Avval ism-familiya va guruhingizni kiriting.", "Сначала укажите имя, фамилию и группу.", "Please enter your full name and group first."),
"name_prompt": _t("Ism-familiya va guruhingizni kiriting:", "Введите имя, фамилию и группу:", "Enter your full name and group:"),
"name_bad": _t("3–150 belgi kiriting.", "Введите от 3 до 150 символов.", "Enter 3–150 characters."),
"name_saved": _t("✅ Saqlandi: <b>{name}</b>", "✅ Сохранено: <b>{name}</b>", "✅ Saved: <b>{name}</b>"),
"b_feedback": _t("✍️ Sharh yozish", "✍️ Написать отзыв", "✍️ Feedback"),
"b_complaint": _t("⚠️ Shikoyat", "⚠️ Жалоба", "⚠️ Complaint"),
"b_news": _t("📰 Yangiliklar", "📰 Новости", "📰 News"),
"b_contest": _t("🏆 Musobaqaga borish", "🏆 Поездка на соревнование", "🏆 Competition trip"),
"b_name": _t("👤 Ism/guruhni o'zgartirish", "👤 Изменить имя/группу", "👤 Change name/group"),
"b_lang": _t("🌐 Til", "🌐 Язык", "🌐 Language"),
"b_cancel": _t("❌ Bekor qilish", "❌ Отмена", "❌ Cancel"),
"write_feedback": _t("Sharh yoki taklifingizni yozing:", "Напишите отзыв или предложение:", "Write your feedback or suggestion:"),
"write_complaint": _t("Shikoyatingizni yozing:", "Опишите жалобу:", "Write your complaint:"),
"short": _t("Xabar juda qisqa.", "Сообщение слишком короткое.", "Message is too short."),
"long": _t("Maksimum {n} belgi.", "Максимум {n} символов.", "Maximum {n} characters."),
"cooldown": _t("{n} soniyadan keyin urinib ko'ring.", "Попробуйте через {n} сек.", "Try again in {n} seconds."),
"sent": _t("✅ Xabaringiz adminga yuborildi.", "✅ Сообщение отправлено администратору.", "✅ Your message was sent to the admin."),
"news_none": _t("Hozircha yangilik yo'q.", "Пока новостей нет.", "No news yet."),
"news_head": _t("📰 <b>So'nggi yangiliklar</b>", "📰 <b>Последние новости</b>", "📰 <b>Latest news</b>"),
"contest_prompt": _t("Musobaqaga qachon borishingizni kiriting.\nFormat: <code>25.10.2026 15:00</code>\nKeyin izoh yozishingiz mumkin, masalan: Toshkent, matematika.", "Укажите дату и время поездки на соревнование.\nФормат: <code>25.10.2026 15:00</code>\nМожно добавить пояснение, например: Ташкент, математика.", "Enter when you plan to go to the competition.\nFormat: <code>25.10.2026 15:00</code>\nYou can add a note, e.g. Tashkent, maths."),
"contest_saved": _t("✅ Сана ва вақт админга юборилди.", "✅ Дата и время отправлены администратору.", "✅ Date and time sent to the admin."),
"admin": _t("🔐 <b>Админ-панель</b>", "🔐 <b>Панель администратора</b>", "🔐 <b>Admin panel</b>"),
"a_all": _t("📋 Barcha yozuvlar", "📋 Все записи", "📋 All records"),
"a_students": _t("👥 O'quvchilar", "👥 Ученики", "👥 Students"),
"a_news": _t("📰 Yangilik boshqaruvi", "📰 Управление новостями", "📰 Manage news"),
"a_export": _t("📤 Eksport", "📤 Экспорт", "📤 Export"),
"a_search": _t("🔎 Qidirish", "🔎 Поиск", "🔎 Search"),
"a_stats": _t("📊 Statistika", "📊 Статистика", "📊 Statistics"),
"a_contests": _t("🏆 Borish rejalari", "🏆 Поездки на соревнования", "🏆 Competition trips"),
"a_cancel": _t("❌ Bekor qilish", "❌ Отмена", "❌ Cancel"),
"search_prompt": _t("Ism, guruh, Telegram ID yoki xabar matnidan qidiring:", "Поиск по имени, группе, Telegram ID или тексту:", "Search by name, group, Telegram ID or message text:"),
"no_results": _t("Natija topilmadi.", "Ничего не найдено.", "No results found."),
"records_head": _t("📋 <b>Oxirgi murojaatlar</b>", "📋 <b>Последние обращения</b>", "📋 <b>Recent messages</b>"),
"students_head": _t("👥 <b>O'quvchilar</b>", "👥 <b>Ученики</b>", "👥 <b>Students</b>"),
"news_admin_prompt": _t("Yangilik matnini yuboring:", "Отправьте текст новости:", "Send the news text:"),
"news_added": _t("✅ Yangilik saqlandi va o'quvchilarga yuborildi. Yetkazildi: {ok}, xato: {fail}", "✅ Новость сохранена и отправлена. Доставлено: {ok}, ошибок: {fail}", "✅ News saved and sent. Delivered: {ok}, failed: {fail}"),
"stats": _t("📊 Sharh: {f}\nShikoyat: {c}\nO'quvchi: {s}\nYangilik: {n}\nMusobaqa rejasi: {p}", "📊 Отзывы: {f}\nЖалобы: {c}\nУченики: {s}\nНовости: {n}\nПоездки: {p}", "📊 Feedback: {f}\nComplaints: {c}\nStudents: {s}\nNews: {n}\nCompetition trips: {p}"),
"delete_confirm": _t("Haqiqatan ham #{id} yozuvni o'chirasizmi?", "Точно удалить запись #{id}?", "Delete record #{id}?"),
"deleted": _t("✅ O'chirildi.", "✅ Удалено.", "✅ Deleted."),
"student_detail": _t("👤 <b>#{id} {name}</b>\nTelegram: <code>{tg}</code>\nUsername: @{username}\n\nO'chirish uchun tugmani bosing.", "👤 <b>#{id} {name}</b>\nTelegram: <code>{tg}</code>\nUsername: @{username}\n\nНажмите кнопку для удаления.", "👤 <b>#{id} {name}</b>\nTelegram: <code>{tg}</code>\nUsername: @{username}\n\nPress the button to delete."),
"student_deleted": _t("O'quvchi va uning yozuvlari o'chirildi.", "Ученик и его записи удалены.", "Student and related records deleted."),
"edit_news_prompt": _t("Tahrirlash uchun yangi matnni yuboring:", "Отправьте новый текст:", "Send the replacement text:"),
"news_saved": _t("✅ Yangilik tahrirlandi.", "✅ Новость обновлена.", "✅ News updated."),
"export_prompt": _t("Eksport turini tanlang:", "Выберите формат экспорта:", "Choose export format:"),
}
def tr(lang, key, **kw):
    return T[key].get(lang, T[key][DEFAULT_LANG]).format(**kw)
def esc(s): return html.escape(s or "")
def now_local(): return datetime.now(TZ).replace(tzinfo=None)
def fmt(dt): return dt.strftime("%d.%m.%Y %H:%M") if dt else "—"
def split_name_group(value):
    value = " ".join((value or "").split())
    if "," in value:
        a, b = value.rsplit(",", 1); return a.strip(), b.strip()
    m = re.search(r"(\d+\s*-?\s*(guruh|группа|group)|(guruh|группа|group)\s*-?\s*\d+)", value, re.I)
    if m: return (value[:m.start()] + value[m.end():]).strip(" ,-") or value, m.group(0)
    return value, ""

class Base(DeclarativeBase): pass
class Student(Base):
    __tablename__ = "students"
    id: Mapped[int] = mapped_column(Integer, primary_key=True, autoincrement=True)
    tg_id: Mapped[int] = mapped_column(BigInteger, unique=True, index=True)
    full_name: Mapped[str] = mapped_column(String(200))
    username: Mapped[Optional[str]] = mapped_column(String(64), nullable=True)
    created_at: Mapped[datetime] = mapped_column(DateTime, default=now_local)
class UserLang(Base):
    __tablename__ = "user_langs"
    tg_id: Mapped[int] = mapped_column(BigInteger, primary_key=True, autoincrement=False)
    lang: Mapped[str] = mapped_column(String(2), default=DEFAULT_LANG)
class Feedback(Base):
    __tablename__ = "feedbacks"
    id: Mapped[int] = mapped_column(Integer, primary_key=True, autoincrement=True)
    student_tg_id: Mapped[int] = mapped_column(BigInteger, index=True)
    student_name: Mapped[str] = mapped_column(String(200))
    username: Mapped[Optional[str]] = mapped_column(String(64), nullable=True)
    text: Mapped[str] = mapped_column(Text)
    kind: Mapped[str] = mapped_column(String(16), default="feedback")
    created_at: Mapped[datetime] = mapped_column(DateTime, default=now_local, index=True)
class News(Base):
    __tablename__ = "news"
    id: Mapped[int] = mapped_column(Integer, primary_key=True, autoincrement=True)
    text: Mapped[str] = mapped_column(Text)
    sent_count: Mapped[int] = mapped_column(Integer, default=0)
    created_at: Mapped[datetime] = mapped_column(DateTime, default=now_local, index=True)
class TripPlan(Base):
    __tablename__ = "competition_trip_plans"
    id: Mapped[int] = mapped_column(Integer, primary_key=True, autoincrement=True)
    student_tg_id: Mapped[int] = mapped_column(BigInteger, index=True)
    student_name: Mapped[str] = mapped_column(String(200))
    username: Mapped[Optional[str]] = mapped_column(String(64), nullable=True)
    departure_at: Mapped[datetime] = mapped_column(DateTime, index=True)
    note: Mapped[str] = mapped_column(Text, default="")
    created_at: Mapped[datetime] = mapped_column(DateTime, default=now_local)
    status: Mapped[str] = mapped_column(String(16), default="planned")

engine = create_async_engine(DATABASE_URL)
Session = async_sessionmaker(engine, expire_on_commit=False)
def _migrate(conn):
    insp = inspect(conn)
    tables = set(insp.get_table_names())
    if "feedbacks" in tables:
        cols = {c["name"] for c in insp.get_columns("feedbacks")}
        if "kind" not in cols:
            conn.execute(sa_text("ALTER TABLE feedbacks ADD COLUMN kind VARCHAR(16) NOT NULL DEFAULT 'feedback'"))
async def init_db():
    async with engine.begin() as conn:
        await conn.run_sync(Base.metadata.create_all)
        await conn.run_sync(_migrate)

async def lang_of(uid):
    async with Session() as s:
        lang = await s.scalar(select(UserLang.lang).where(UserLang.tg_id == uid))
    return lang or DEFAULT_LANG
async def set_lang(uid, lang):
    async with Session() as s:
        row = await s.get(UserLang, uid)
        if row: row.lang = lang
        else: s.add(UserLang(tg_id=uid, lang=lang))
        await s.commit()
async def get_student(uid):
    async with Session() as s: return await s.scalar(select(Student).where(Student.tg_id == uid))
async def save_student(uid, name, username):
    async with Session() as s:
        row = await s.scalar(select(Student).where(Student.tg_id == uid))
        if row: row.full_name, row.username = name, username
        else: row = Student(tg_id=uid, full_name=name, username=username); s.add(row)
        await s.commit()
        return row
async def get_recipients():
    async with Session() as s:
        ids = list((await s.scalars(select(Student.tg_id))).all())
    return [i for i in ids if i not in ADMIN_IDS]
async def add_feedback(student, username, body, kind):
    async with Session() as s:
        row = Feedback(student_tg_id=student.tg_id, student_name=student.full_name, username=username, text=body, kind=kind, created_at=now_local())
        s.add(row); await s.commit(); return row.id
async def recent_feedback(limit=30):
    async with Session() as s: return list((await s.scalars(select(Feedback).order_by(desc(Feedback.id)).limit(limit))).all())
async def all_feedback():
    async with Session() as s: return list((await s.scalars(select(Feedback).order_by(Feedback.id.asc()))).all())
async def search_all(term):
    like = f"%{term}%"
    async with Session() as s:
        fb = list((await s.scalars(select(Feedback).where(
            (Feedback.student_name.ilike(like)) | (Feedback.text.ilike(like)) |
            (Feedback.username.ilike(like)) | (Feedback.student_tg_id.cast(String).ilike(like))
        ).order_by(desc(Feedback.id)).limit(20))).all())
        students = list((await s.scalars(select(Student).where(
            (Student.full_name.ilike(like)) | (Student.username.ilike(like)) |
            (Student.tg_id.cast(String).ilike(like))
        ).order_by(desc(Student.id)).limit(15))).all())
    return fb, students
async def all_students():
    async with Session() as s: return list((await s.scalars(select(Student).order_by(desc(Student.id)).limit(50))).all())
async def delete_feedback(fid):
    async with Session() as s:
        row = await s.get(Feedback, fid)
        if not row: return False
        await s.delete(row); await s.commit(); return True
async def delete_student(sid):
    async with Session() as s:
        row = await s.get(Student, sid)
        if not row: return False
        tg = row.tg_id
        await s.execute(sa_text("DELETE FROM feedbacks WHERE student_tg_id=:tg"), {"tg": tg})
        await s.execute(sa_text("DELETE FROM competition_trip_plans WHERE student_tg_id=:tg"), {"tg": tg})
        await s.execute(sa_text("DELETE FROM user_langs WHERE tg_id=:tg"), {"tg": tg})
        await s.delete(row); await s.commit(); return True
async def add_news(body, sent):
    async with Session() as s:
        row = News(text=body, sent_count=sent, created_at=now_local()); s.add(row); await s.commit(); return row.id
async def get_news_list():
    async with Session() as s: return list((await s.scalars(select(News).order_by(desc(News.id)).limit(20))).all())
async def edit_news(nid, body):
    async with Session() as s:
        row = await s.get(News, nid)
        if not row: return False
        row.text = body; await s.commit(); return True
async def save_trip(student, username, departure_at, note):
    async with Session() as s:
        row = TripPlan(student_tg_id=student.tg_id, student_name=student.full_name, username=username, departure_at=departure_at, note=note, created_at=now_local())
        s.add(row); await s.commit(); return row.id
async def trip_list():
    async with Session() as s: return list((await s.scalars(select(TripPlan).order_by(desc(TripPlan.departure_at)).limit(100))).all())
async def stats():
    async with Session() as s:
        f = await s.scalar(select(func.count(Feedback.id)).where(Feedback.kind == "feedback")) or 0
        c = await s.scalar(select(func.count(Feedback.id)).where(Feedback.kind == "complaint")) or 0
        st = await s.scalar(select(func.count(Student.id))) or 0
        n = await s.scalar(select(func.count(News.id))) or 0
        p = await s.scalar(select(func.count(TripPlan.id))) or 0
    return f,c,st,n,p

def make_csv(filename, headers, rows):
    b = io.StringIO(newline=""); w = csv.writer(b, delimiter=";"); w.writerow(headers); w.writerows(rows)
    return BufferedInputFile(b.getvalue().encode("utf-8-sig"), filename=filename)
def make_xlsx(filename, headers, rows):
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill, Alignment
    wb = Workbook(); ws = wb.active; ws.title = "Export"
    ws.append(headers)
    for c in ws[1]:
        c.font = Font(bold=True, color="FFFFFF")
        c.fill = PatternFill("solid", fgColor="2F5597")
        c.alignment = Alignment(wrap_text=True)
    for row in rows:
        ws.append([("\'" + str(v)) if isinstance(v, str) and v[:1] in ("=", "+", "-", "@") else (str(v) if v is not None else "") for v in row])
    ws.freeze_panes = "A2"; ws.auto_filter.ref = ws.dimensions
    for col in ws.columns:
        letter = col[0].column_letter
        width = min(max(max(len(str(c.value or "")) for c in col) + 2, 12), 48)
        ws.column_dimensions[letter].width = width
    b = io.BytesIO(); wb.save(b)
    return BufferedInputFile(b.getvalue(), filename=filename)
def make_pdf(filename, title, headers, rows):
    from reportlab.lib import colors
    from reportlab.lib.pagesizes import landscape, A4
    from reportlab.lib.styles import getSampleStyleSheet
    from reportlab.lib.units import mm
    from reportlab.platypus import SimpleDocTemplate, Table, TableStyle, Paragraph, Spacer
    from reportlab.pdfbase import pdfmetrics
    from reportlab.pdfbase.ttfonts import TTFont
    # Unicode font if DejaVu is available; otherwise built-in font is used.
    font_name = "Helvetica"
    for path in ("/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf", r"C:\Windows\Fonts\arial.ttf"):
        if os.path.exists(path):
            try:
                pdfmetrics.registerFont(TTFont("ExportUnicode", path)); font_name = "ExportUnicode"; break
            except Exception: pass
    styles = getSampleStyleSheet()
    styles["Title"].fontName = font_name
    styles["Normal"].fontName = font_name
    b = io.BytesIO()
    doc = SimpleDocTemplate(b, pagesize=landscape(A4), rightMargin=10*mm, leftMargin=10*mm, topMargin=10*mm, bottomMargin=10*mm)
    data = [[Paragraph(esc(str(x)), styles["Normal"]) for x in headers]]
    for row in rows:
        data.append([Paragraph(esc(str(x if x is not None else ""))[:1500], styles["Normal"]) for x in row])
    table = Table(data, repeatRows=1)
    table.setStyle(TableStyle([
        ("BACKGROUND",(0,0),(-1,0),colors.HexColor("#2F5597")),
        ("TEXTCOLOR",(0,0),(-1,0),colors.white),
        ("GRID",(0,0),(-1,-1),0.35,colors.grey),
        ("VALIGN",(0,0),(-1,-1),"TOP"),
        ("FONTNAME",(0,0),(-1,-1),font_name),
        ("FONTSIZE",(0,0),(-1,-1),7),
        ("LEFTPADDING",(0,0),(-1,-1),4), ("RIGHTPADDING",(0,0),(-1,-1),4),
    ]))
    doc.build([Paragraph(esc(title), styles["Title"]), Spacer(1, 5*mm), table])
    return BufferedInputFile(b.getvalue(), filename=filename)

def lang_kb():
    b=InlineKeyboardBuilder()
    for k,v in LANGS.items(): b.button(text=v, callback_data=f"lang:{k}")
    b.adjust(1); return b.as_markup()
def student_menu(lang):
    return ReplyKeyboardMarkup(keyboard=[
        [KeyboardButton(text=tr(lang,"b_feedback")), KeyboardButton(text=tr(lang,"b_complaint"))],
        [KeyboardButton(text=tr(lang,"b_news")), KeyboardButton(text=tr(lang,"b_contest"))],
        [KeyboardButton(text=tr(lang,"b_name")), KeyboardButton(text=tr(lang,"b_lang"))],
    ], resize_keyboard=True)
def admin_menu(lang):
    return ReplyKeyboardMarkup(keyboard=[
        [KeyboardButton(text=tr(lang,"a_all")), KeyboardButton(text=tr(lang,"a_search"))],
        [KeyboardButton(text=tr(lang,"a_students")), KeyboardButton(text=tr(lang,"a_contests"))],
        [KeyboardButton(text=tr(lang,"a_news")), KeyboardButton(text=tr(lang,"a_export"))],
        [KeyboardButton(text=tr(lang,"a_stats")), KeyboardButton(text=tr(lang,"b_lang"))],
    ], resize_keyboard=True)
def cancel_inline():
    return InlineKeyboardMarkup(inline_keyboard=[[InlineKeyboardButton(text="❌ Cancel", callback_data="cancel")]])
def confirm_delete(fid):
    return InlineKeyboardMarkup(inline_keyboard=[[InlineKeyboardButton(text="🗑 Delete", callback_data=f"del:{fid}")],[InlineKeyboardButton(text="Cancel",callback_data="cancel")]])
def admin_feedback_kb(fid):
    return InlineKeyboardMarkup(inline_keyboard=[[InlineKeyboardButton(text=f"🗑 Delete #{fid}",callback_data=f"askdel:{fid}")]])
def export_kb():
    return InlineKeyboardMarkup(inline_keyboard=[
        [InlineKeyboardButton(text="📊 Excel (.xlsx)",callback_data="export:xlsx")],
        [InlineKeyboardButton(text="📄 PDF",callback_data="export:pdf")],
        [InlineKeyboardButton(text="📑 CSV",callback_data="export:csv")],
    ])

common, student_router, admin_router = Router(), Router(), Router()
admin_router.message.filter(F.from_user.id.in_(ADMIN_IDS))
admin_router.callback_query.filter(F.from_user.id.in_(ADMIN_IDS))
student_router.message.filter(~F.from_user.id.in_(ADMIN_IDS))
student_router.callback_query.filter(~F.from_user.id.in_(ADMIN_IDS))
class Reg(StatesGroup): name=State()
class Compose(StatesGroup): feedback=State(); complaint=State()
class Trip(StatesGroup): date=State(); note=State()
class AdminState(StatesGroup): search=State(); news=State(); edit_news=State()

@common.message(Command("id"))
async def get_id(m):
    await m.answer(f"Telegram ID: <code>{m.from_user.id}</code>")
@common.message(Command("lang"))
@common.message(F.text.in_({tr(l,"b_lang") for l in LANGS}))
async def choose_lang(m):
    await m.answer(tr(DEFAULT_LANG,"choose_lang"), reply_markup=lang_kb())
@common.callback_query(F.data.startswith("lang:"))
async def lang_change(c, state: FSMContext):
    lang=(c.data or "").split(":")[-1]
    if lang not in LANGS: await c.answer(); return
    await set_lang(c.from_user.id,lang); await state.clear(); await c.answer()
    if c.message:
        try: await c.message.delete()
        except Exception: pass
    if c.from_user.id in ADMIN_IDS:
        await c.bot.send_message(c.from_user.id,tr(lang,"admin"),reply_markup=admin_menu(lang))
    else:
        st=await get_student(c.from_user.id)
        if st: await c.bot.send_message(c.from_user.id,tr(lang,"hello",name=esc(st.full_name)),reply_markup=student_menu(lang))
        else:
            await state.set_state(Reg.name)
            await c.bot.send_message(c.from_user.id,tr(lang,"welcome"),reply_markup=ReplyKeyboardRemove())
@student_router.message(CommandStart())
async def start_student(m,state: FSMContext):
    await state.clear(); lang=await lang_of(m.from_user.id)
    if lang not in LANGS: lang=DEFAULT_LANG
    st=await get_student(m.from_user.id)
    if st: await m.answer(tr(lang,"hello",name=esc(st.full_name)),reply_markup=student_menu(lang))
    else:
        await state.set_state(Reg.name); await m.answer(tr(lang,"welcome"),reply_markup=ReplyKeyboardRemove())
@admin_router.message(CommandStart())
@admin_router.message(Command("admin"))
async def start_admin(m,state: FSMContext):
    await state.clear(); lang=await lang_of(m.from_user.id)
    await m.answer(tr(lang,"admin"),reply_markup=admin_menu(lang))
@student_router.message(F.text.in_({tr(l,"b_cancel") for l in LANGS}))
async def cancel_student(m,state):
    await state.clear(); lang=await lang_of(m.from_user.id)
    await m.answer(tr(lang,"cancelled"),reply_markup=student_menu(lang))
@student_router.message(Command("name"))
@student_router.message(F.text.in_({tr(l,"b_name") for l in LANGS}))
async def rename(m,state):
    lang=await lang_of(m.from_user.id); await state.set_state(Reg.name)
    await m.answer(tr(lang,"name_prompt"),reply_markup=ReplyKeyboardRemove())
@student_router.message(Reg.name,F.text,~F.text.startswith("/"))
async def register_name(m,state):
    lang=await lang_of(m.from_user.id); name=" ".join(m.text.split())
    if not 3<=len(name)<=150: await m.answer(tr(lang,"name_bad")); return
    await save_student(m.from_user.id,name,m.from_user.username); await state.clear()
    await m.answer(tr(lang,"name_saved",name=esc(name)),reply_markup=student_menu(lang))
async def compose_start(m,state,kind):
    lang=await lang_of(m.from_user.id)
    if not await get_student(m.from_user.id):
        await state.set_state(Reg.name); await m.answer(tr(lang,"need_name")); return
    await state.set_state(Compose.feedback if kind=="feedback" else Compose.complaint)
    await m.answer(tr(lang,"write_feedback" if kind=="feedback" else "write_complaint"),reply_markup=ReplyKeyboardMarkup(keyboard=[[KeyboardButton(text=tr(lang,"b_cancel"))]],resize_keyboard=True))
@student_router.message(F.text.in_({tr(l,"b_feedback") for l in LANGS}))
async def feedback_start(m,state): await compose_start(m,state,"feedback")
@student_router.message(F.text.in_({tr(l,"b_complaint") for l in LANGS}))
async def complaint_start(m,state): await compose_start(m,state,"complaint")
async def submit(m,state,kind):
    lang=await lang_of(m.from_user.id); st=await get_student(m.from_user.id)
    if not st: await state.set_state(Reg.name); await m.answer(tr(lang,"need_name")); return
    body=(m.text or "").strip()
    if len(body)<MIN_LEN: await m.answer(tr(lang,"short")); return
    if len(body)>MAX_LEN: await m.answer(tr(lang,"long",n=MAX_LEN)); return
    now=time.monotonic(); wait=COOLDOWN_SEC-(now-_last_sent.get(m.from_user.id,0))
    if wait>0: await m.answer(tr(lang,"cooldown",n=int(wait)+1)); return
    _last_sent[m.from_user.id]=now
    fid=await add_feedback(st,m.from_user.username,body,kind); await state.clear()
    await m.answer(tr(lang,"sent"),reply_markup=student_menu(lang))
    for aid in ADMIN_IDS:
        try:
            await m.bot.send_message(aid,f"{'📝 Sharh' if kind=='feedback' else '⚠️ Shikoyat'} #{fid}\n👤 {esc(st.full_name)}\n🆔 <code>{st.tg_id}</code>\n🕒 {fmt(now_local())}\n\n{esc(body)}",reply_markup=admin_feedback_kb(fid),protect_content=True)
        except Exception: logging.exception("Admin notification failed: %s",aid)
@student_router.message(Compose.feedback,F.text,~F.text.startswith("/"))
async def send_feedback(m,state): await submit(m,state,"feedback")
@student_router.message(Compose.complaint,F.text,~F.text.startswith("/"))
async def send_complaint(m,state): await submit(m,state,"complaint")
@student_router.message(F.text.in_({tr(l,"b_news") for l in LANGS}))
async def student_news(m,state):
    await state.clear(); lang=await lang_of(m.from_user.id); items=await get_news_list()
    if not items: await m.answer(tr(lang,"news_none"),reply_markup=student_menu(lang)); return
    await m.answer(tr(lang,"news_head"),reply_markup=student_menu(lang))
    for item in reversed(items): await m.answer(f"🕒 {fmt(item.created_at)}\n\n{item.text}",protect_content=True)
@student_router.message(F.text.in_({tr(l,"b_contest") for l in LANGS}))
async def trip_start(m,state):
    lang=await lang_of(m.from_user.id)
    if not await get_student(m.from_user.id): await state.set_state(Reg.name); await m.answer(tr(lang,"need_name")); return
    await state.set_state(Trip.date); await m.answer(tr(lang,"contest_prompt"),reply_markup=ReplyKeyboardMarkup(keyboard=[[KeyboardButton(text=tr(lang,"b_cancel"))]],resize_keyboard=True))
@student_router.message(Trip.date,F.text,~F.text.startswith("/"))
async def trip_date(m,state):
    lang=await lang_of(m.from_user.id); raw=" ".join(m.text.split())
    match=re.fullmatch(r"(\d{1,2})[./](\d{1,2})[./](\d{4})\s+(\d{1,2}):(\d{2})(?:\s*[|,]\s*(.*))?",raw)
    if not match:
        await m.answer("Format: 25.10.2026 15:00 yoki 25.10.2026 15:00, Toshkent")
        return
    d,mo,y,h,mi=map(int,match.groups()[:5]); note=(match.group(6) or "").strip()
    try: dt=datetime(y,mo,d,h,mi)
    except ValueError: await m.answer("Sana yoki vaqt noto'g'ri. Qayta kiriting."); return
    if dt < now_local(): await m.answer("Kelajakdagi sana va vaqtni kiriting."); return
    await state.update_data(departure=dt.isoformat(sep=" "),note=note)
    if note:
        await finalize_trip(m,state,dt,note); return
    await state.set_state(Trip.note)
    await m.answer("Izohni yozing yoki '-' yuboring (ixtiyoriy):")
async def finalize_trip(m,state,dt,note):
    lang=await lang_of(m.from_user.id); st=await get_student(m.from_user.id)
    if not st: await state.clear(); await m.answer(tr(lang,"need_name")); return
    tid=await save_trip(st,m.from_user.username,dt,note)
    await state.clear(); await m.answer(tr(lang,"contest_saved"),reply_markup=student_menu(lang))
    for aid in ADMIN_IDS:
        try: await m.bot.send_message(aid,f"🏆 <b>Yangi musobaqaga borish rejasi #{tid}</b>\n👤 {esc(st.full_name)}\n🆔 <code>{st.tg_id}</code>\n📅 {fmt(dt)}\n📝 {esc(note or '—')}",protect_content=True)
        except Exception: logging.exception("Trip admin notification failed")
@student_router.message(Trip.note,F.text,~F.text.startswith("/"))
async def trip_note(m,state):
    data=await state.get_data(); dt=datetime.fromisoformat(data["departure"]); note="" if m.text.strip()=="-" else m.text.strip()
    if len(note)>500: await m.answer("Izoh 500 belgidan oshmasin."); return
    await finalize_trip(m,state,dt,note)

async def send_export(m, fmt_name):
    fbs=await all_feedback()
    headers=["ID","Type","Date","Full name","Group","Telegram ID","Username","Text"]
    rows=[]
    for x in fbs:
        name,group=split_name_group(x.student_name)
        rows.append([x.id,x.kind,fmt(x.created_at),name,group,x.student_tg_id,x.username or "",x.text])
    if fmt_name=="xlsx": doc=make_xlsx("feedback_export.xlsx",headers,rows)
    elif fmt_name=="pdf": doc=make_pdf("feedback_export.pdf","Feedback and complaints",headers,rows)
    else: doc=make_csv("feedback_export.csv",headers,rows)
    await m.answer_document(doc,caption=f"Feedback/complaints: {len(rows)}",protect_content=True)
@admin_router.message(F.text.in_({tr(l,"a_all") for l in LANGS}))
async def admin_all(m):
    lang=await lang_of(m.from_user.id); rows=await recent_feedback(25)
    if not rows: await m.answer(tr(lang,"no_results")); return
    for x in rows:
        await m.answer(f"#{x.id} {'📝' if x.kind=='feedback' else '⚠️'} {esc(x.student_name)}\n🆔 <code>{x.student_tg_id}</code> · {fmt(x.created_at)}\n{esc(x.text[:700])}",reply_markup=admin_feedback_kb(x.id),protect_content=True)
@admin_router.callback_query(F.data.startswith("askdel:"))
async def ask_delete(c):
    fid=int(c.data.split(":")[1]); lang=await lang_of(c.from_user.id)
    await c.message.answer(tr(lang,"delete_confirm",id=fid),reply_markup=confirm_delete(fid)); await c.answer()
@admin_router.callback_query(F.data.startswith("del:"))
async def do_delete(c):
    fid=int(c.data.split(":")[1]); ok=await delete_feedback(fid); lang=await lang_of(c.from_user.id)
    await c.answer(tr(lang,"deleted") if ok else tr(lang,"no_results"),show_alert=True)
    if c.message:
        try: await c.message.edit_reply_markup(reply_markup=None)
        except Exception: pass
@admin_router.callback_query(F.data=="cancel")
async def cb_cancel(c,state):
    await state.clear(); await c.answer("Cancelled")
    if c.message:
        try: await c.message.edit_reply_markup(reply_markup=None)
        except Exception: pass
@admin_router.message(F.text.in_({tr(l,"a_search") for l in LANGS}))
async def search_start(m,state):
    lang=await lang_of(m.from_user.id); await state.set_state(AdminState.search); await m.answer(tr(lang,"search_prompt"),reply_markup=cancel_inline())
@admin_router.message(AdminState.search,F.text)
async def search_run(m,state):
    lang=await lang_of(m.from_user.id); term=m.text.strip()
    if len(term)<2: await m.answer("Kamida 2 ta belgi kiriting."); return
    fb, students=await search_all(term); await state.clear()
    if not fb and not students: await m.answer(tr(lang,"no_results")); return
    for x in fb:
        await m.answer(f"#{x.id} {'📝' if x.kind=='feedback' else '⚠️'} {esc(x.student_name)}\n🆔 <code>{x.student_tg_id}</code> · {fmt(x.created_at)}\n{esc(x.text[:500])}",reply_markup=admin_feedback_kb(x.id),protect_content=True)
    for x in students:
        await m.answer(tr(lang,"student_detail",id=x.id,name=esc(x.full_name),tg=x.tg_id,username=esc(x.username or "—")),reply_markup=InlineKeyboardMarkup(inline_keyboard=[[InlineKeyboardButton(text=f"🗑 Delete student #{x.id}",callback_data=f"studentdel:{x.id}")]]))
@admin_router.message(F.text.in_({tr(l,"a_students") for l in LANGS}))
async def show_students(m):
    lang=await lang_of(m.from_user.id); rows=await all_students()
    if not rows: await m.answer(tr(lang,"no_results")); return
    await m.answer(tr(lang,"students_head"))
    for x in rows:
        await m.answer(tr(lang,"student_detail",id=x.id,name=esc(x.full_name),tg=x.tg_id,username=esc(x.username or "—")),reply_markup=InlineKeyboardMarkup(inline_keyboard=[[InlineKeyboardButton(text=f"🗑 Delete student #{x.id}",callback_data=f"studentdel:{x.id}")]]),protect_content=True)
@admin_router.callback_query(F.data.startswith("studentdel:"))
async def student_delete_cb(c):
    sid=int(c.data.split(":")[1]); ok=await delete_student(sid); lang=await lang_of(c.from_user.id)
    await c.answer(tr(lang,"student_deleted") if ok else tr(lang,"no_results"),show_alert=True)
    if c.message:
        try: await c.message.edit_reply_markup(reply_markup=None)
        except Exception: pass
@admin_router.message(F.text.in_({tr(l,"a_news") for l in LANGS}))
async def news_admin(m,state):
    lang=await lang_of(m.from_user.id); items=await get_news_list()
    b=InlineKeyboardBuilder()
    for n in items: b.button(text=f"✏️ #{n.id} {n.text[:30]}",callback_data=f"editnews:{n.id}")
    b.adjust(1); b.row(InlineKeyboardButton(text="➕ New news",callback_data="newnews"))
    await m.answer("📰 News management",reply_markup=b.as_markup())
@admin_router.callback_query(F.data=="newnews")
async def new_news(c,state):
    lang=await lang_of(c.from_user.id); await state.set_state(AdminState.news); await c.answer()
    await c.message.answer(tr(lang,"news_admin_prompt"),reply_markup=cancel_inline())
@admin_router.callback_query(F.data.startswith("editnews:"))
async def edit_news_start(c,state):
    nid=int(c.data.split(":")[1]); await state.update_data(news_id=nid); await state.set_state(AdminState.edit_news); lang=await lang_of(c.from_user.id); await c.answer()
    await c.message.answer(tr(lang,"edit_news_prompt"),reply_markup=cancel_inline())
@admin_router.message(AdminState.edit_news,F.text)
async def edit_news_save(m,state):
    data=await state.get_data(); lang=await lang_of(m.from_user.id); ok=await edit_news(data["news_id"],m.text.strip()); await state.clear()
    await m.answer(tr(lang,"news_saved") if ok else tr(lang,"no_results"))
@admin_router.message(AdminState.news,F.text)
async def new_news_save(m,state):
    lang=await lang_of(m.from_user.id); body=m.html_text
    if not body.strip(): await m.answer(tr(lang,"news_admin_prompt")); return
    ok=fail=0
    for uid in await get_recipients():
        try: await m.bot.send_message(uid,body,protect_content=True); ok+=1
        except TelegramRetryAfter as e:
            await asyncio.sleep(e.retry_after+1)
            try: await m.bot.send_message(uid,body,protect_content=True); ok+=1
            except Exception: fail+=1
        except TelegramForbiddenError: fail+=1
        except Exception: fail+=1
        await asyncio.sleep(.04)
    await add_news(body,ok); await state.clear()
    await m.answer(tr(lang,"news_added",ok=ok,fail=fail),reply_markup=admin_menu(lang))
@admin_router.message(F.text.in_({tr(l,"a_export") for l in LANGS}))
async def export_start(m):
    lang=await lang_of(m.from_user.id); await m.answer(tr(lang,"export_prompt"),reply_markup=export_kb())
@admin_router.callback_query(F.data.startswith("export:"))
async def export_cb(c):
    fmt_name=c.data.split(":")[1]; lang=await lang_of(c.from_user.id)
    # Generate current full export; export functions are synchronous and only process in-memory records.
    rows=await all_feedback()
    headers=["ID","Type","Date","Full name","Group","Telegram ID","Username","Text"]
    data=[]
    for x in rows:
        name,group=split_name_group(x.student_name)
        data.append([x.id,x.kind,fmt(x.created_at),name,group,x.student_tg_id,x.username or "",x.text])
    if fmt_name=="xlsx": doc=make_xlsx("feedback_export.xlsx",headers,data)
    elif fmt_name=="pdf": doc=make_pdf("feedback_export.pdf","Feedback and complaints",headers,data)
    elif fmt_name=="csv": doc=make_csv("feedback_export.csv",headers,data)
    else: await c.answer("Unknown format",show_alert=True); return
    await c.message.answer_document(doc,caption=f"Export: {len(data)} records",protect_content=True); await c.answer()
@admin_router.message(F.text.in_({tr(l,"a_stats") for l in LANGS}))
async def stats_admin(m):
    lang=await lang_of(m.from_user.id); f,c,s,n,p=await stats(); await m.answer(tr(lang,"stats",f=f,c=c,s=s,n=n,p=p))
@admin_router.message(F.text.in_({tr(l,"a_contests") for l in LANGS}))
async def trips_admin(m):
    lang=await lang_of(m.from_user.id); rows=await trip_list()
    if not rows: await m.answer(tr(lang,"no_results")); return
    for x in rows:
        await m.answer(f"🏆 <b>#{x.id}</b> · {fmt(x.departure_at)}\n👤 {esc(x.student_name)}\n🆔 <code>{x.student_tg_id}</code>\nUsername: @{esc(x.username or '—')}\n📝 {esc(x.note or '—')}",reply_markup=InlineKeyboardMarkup(inline_keyboard=[[InlineKeyboardButton(text="🗑 Delete",callback_data=f"tripdel:{x.id}")]]),protect_content=True)
@admin_router.callback_query(F.data.startswith("tripdel:"))
async def trip_delete(c):
    tid=int(c.data.split(":")[1])
    async with Session() as s:
        row=await s.get(TripPlan,tid)
        if row: await s.delete(row); await s.commit()
    await c.answer("Deleted" if row else "Not found",show_alert=True)
    if c.message:
        try: await c.message.edit_reply_markup(reply_markup=None)
        except Exception: pass
@admin_router.message()
async def admin_other(m):
    if m.text: return
    await m.answer("Please send text.")

@student_router.message()
async def student_other(m):
    lang=await lang_of(m.from_user.id); await m.answer(tr(lang,"only_text"))

async def main():
    logging.basicConfig(level=logging.INFO,format="%(asctime)s | %(levelname)s | %(message)s")
    if not TOKEN: raise SystemExit("BOT_TOKEN .env faylida ko'rsatilmagan.")
    if not ADMIN_IDS: raise SystemExit("ADMIN_IDS .env faylida ko'rsatilmagan.")
    await init_db()
    bot=Bot(TOKEN,default=DefaultBotProperties(parse_mode=ParseMode.HTML))
    dp=Dispatcher(storage=MemoryStorage())
    dp.include_router(common); dp.include_router(admin_router); dp.include_router(student_router)
    try:
        await dp.start_polling(bot)
    finally:
        await engine.dispose()
if __name__=="__main__":
    try: asyncio.run(main())
    except KeyboardInterrupt: print("Bot stopped.")
